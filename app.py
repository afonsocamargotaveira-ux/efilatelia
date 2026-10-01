"""
Agregador de leilões de selos (somente leilões EM ANDAMENTO)

Executar:
    pip install streamlit requests beautifulsoup4 pandas
    streamlit run app.py

Fonte principal: portal LeilõesBR, endpoint "busca_andamento.asp" (busca apenas em
leilões em andamento). O portal reúne várias casas filatélicas (Filatélica BH,
Galeria Filatélica e Numismática, Cida Mello, Coins e Stamps, Moeda e Selo,
Universo das Coleções, Inovar, etc.).
"""

from __future__ import annotations

import base64
import html
import io
import json
import re
import time
import unicodedata
from datetime import datetime
from concurrent.futures import ThreadPoolExecutor, as_completed
from concurrent.futures import TimeoutError as FutTimeout
from dataclasses import dataclass
from dataclasses import asdict, replace
from pathlib import Path
from urllib.parse import quote, quote_plus, urljoin, urlparse
from urllib.robotparser import RobotFileParser

import pandas as pd
import requests
import streamlit as st
from bs4 import BeautifulSoup

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
    ),
    "Accept-Language": "pt-BR,pt;q=0.9,en;q=0.8",
}
TIMEOUT = (6, 15)  # (conexão, leitura) em segundos

ENCERRADO = re.compile(
    r"\b(encerrad[oa]|finalizad[oa]|vendid[oa]|arrematad[oa]|ended|closed|sold)\b",
    re.IGNORECASE,
)
# Indícios, na página do lote, de que o leilão já acabou
ENCERRADO_NA_PAGINA = re.compile(
    r"(lote vendido|lote encerrado|leil[aã]o (encerrado|finalizado)|"
    r"venda p[óo]s preg[aã]o|auction (has )?ended)",
    re.IGNORECASE,
)


@dataclass
class Site:
    nome: str
    url_busca: str            # use {q} para o termo e (opcional) {pag} para a página
    seletor_link: str         # seletor CSS dos <a> que apontam para cada lote
    encoding: str = "utf-8"   # codificação usada na query string
    usa_paginas: bool = False # True se a URL tiver {pag}
    grupo: str = ""          # nome exibido nos resultados (agrupa várias categorias)
    busca_local: bool = False # True: o site não filtra por termo; baixamos as páginas e filtramos aqui
    min_paginas: int = 0      # para busca_local: lê pelo menos este nº de páginas do catálogo
    sitemap: str = ""         # sitemap.xml público do site (lista de produtos), usado antes do catálogo
    param_pagina: str = ""    # acrescentado à URL só a partir da página 2 (ex.: "&pag={pag}"); a página 1 fica idêntica à URL do site
    url_reserva: str = ""     # catálogo paginado usado se a URL principal for barrada pelo robots.txt
    filtro_ativos: str = ""   # filtro especial de "leilão em andamento" (ex.: "philasearch")
    sessao: bool = False      # True: visita a home antes (cookies) e envia Referer
    js: bool = False          # True: a página precisa de JavaScript (usa Playwright, se ativado)
    extra: bool = False       # True: site adicionado pelo usuário (salvo em sites_extras.json)
    pag_1_vazia: bool = False # True: na 1ª página, {pag} vira texto vazio (ex.: "page=")
    href_exige: str = ""      # regex: só links cujo endereço combine (ex.: páginas de produto)
    href_exclui: str = ""     # regex: descarta links de menu, categorias, carrinho etc.
    ativo: bool = True


SITES_PADRAO = [
    # Delcampe: configuração mantida como estava (funciona bem). O flag js=False só garante
    # que ele continue sendo lido pelo mesmo caminho mesmo se o Playwright for ligado.
    Site(
        nome="Delcampe – selos",
        grupo="Delcampe – selos",
        url_busca="https://www.delcampe.net/en_US/collectibles/search?term={q}&page={pag}",
        seletor_link='a[href*="/collectibles/"][href$=".html"]',
        usa_paginas=True,
        js=False,
    ),
    Site(
        nome="HipStamp – leilões de selos",
        grupo="HipStamp – leilões de selos",
        # URL exata: https://www.hipstamp.com/search?keywords=guiana&category_id=
        url_busca="https://www.hipstamp.com/search?keywords={q}&category_id=",
        param_pagina="&page={pag}",
        seletor_link='a[href*="/listing/"]',
        usa_paginas=True,
    ),
    # ---- Demais sites: a URL exata de busca de cada site, com o termo trocado por {q}.
    # ---- A página 1 é idêntica à URL do site; só as páginas 2+ recebem o parâmetro extra.
    Site(
        nome="LeilõesBR (leilões em andamento)",
        grupo="LeilõesBR (leilões em andamento)",
        # URL exata: https://www.leiloesbr.com.br/busca_andamento.asp?op=2&b=&ga=*&uf=*&pesquisa=guiana
        url_busca="https://www.leiloesbr.com.br/busca_andamento.asp?op=2&b=&ga=*&uf=*&pesquisa={q}",
        param_pagina="&pag={pag}",
        seletor_link='a[href*="abre_catalogo.asp?t=1"]',
        encoding="latin-1",  # o portal usa ISO-8859-1
        usa_paginas=True,
    ),
    Site(
        nome="Philasearch (leilões internacionais)",
        grupo="Philasearch (leilões internacionais)",
        # URL exata: ...lots?page=&set_gesetz_bestaetigt_jn=J&gesetz_bestaetigt_neu=J&set_sprache=en&suchtext=guiana
        url_busca=(
            "https://www.philasearch.com/en/lots?page={pag}&set_gesetz_bestaetigt_jn=J"
            "&gesetz_bestaetigt_neu=J&set_sprache=en&suchtext={q}"
        ),
        seletor_link='a[href*="posdetail"], a[href*="info.php3?losnr"], a[href*="/i_"]',
        usa_paginas=True,
        filtro_ativos="philasearch",
        sessao=True,
        pag_1_vazia=True,  # na 1ª página o original usa "page=" vazio
        js=True,           # precisa do Playwright (o site recusa requisições comuns)
    ),
    Site(
        nome="Invaluable (leilões de selos)",
        grupo="Invaluable (leilões de selos)",
        # URL exata: https://www.invaluable.com/search?query=guiana&keyword=guiana
        url_busca="https://www.invaluable.com/search?query={q}&keyword={q}",
        param_pagina="&page={pag}",
        seletor_link='a[href*="/auction-lot/"]',
        usa_paginas=True,
        js=True,           # precisa do Playwright
    ),
    Site(
        nome="Loja de Selos (venda direta, NÃO é leilão)",
        grupo="Loja de Selos (venda direta, NÃO é leilão)",
        # URL exata: https://www.lojadeselos.com.br/search/?q=guiana
        # Se o robots.txt do site barrar /search/, o app usa o sitemap e, depois, o catálogo.
        url_busca="https://www.lojadeselos.com.br/search/?q={q}",
        param_pagina="&mpage={pag}",
        url_reserva="https://www.lojadeselos.com.br/produtos/?mpage={pag}",
        sitemap="https://www.lojadeselos.com.br/sitemap.xml",
        seletor_link='a[href*="produto"]',
        usa_paginas=True,
        busca_local=True,
        min_paginas=12,
        href_exige=r"/produtos?/[^/?#]+",
        href_exclui=r"(mpage=|categoria|carrinho|checkout|login|conta|contato|sobre|"
                    r"politica|termos|blog|whatsapp|mailto:|tel:)",
    ),
    Site(
        nome="Olho de Boi Filatelia (venda direta, NÃO é leilão)",
        grupo="Olho de Boi Filatelia (venda direta, NÃO é leilão)",
        # URL exata: https://olhodeboifilatelia.com.br/loja-OlhodeBoiFilatelia-_encontre_palavras_-guiana
        url_busca="https://olhodeboifilatelia.com.br/loja-OlhodeBoiFilatelia-_encontre_palavras_-{qp}",
        # O nome do parâmetro de página não está confirmado; se o site o ignorar, o app
        # percebe a repetição e para na página 1.
        param_pagina="-_pagina_-{pag}",
        seletor_link='a[href*="/oferta-"]',
        encoding="latin-1",  # a página declara ISO-8859-1
        usa_paginas=True,
    ),
]


# ----------------------------------------------------------------------------
# Utilidades
# ----------------------------------------------------------------------------
def norm(s: str) -> str:
    s = unicodedata.normalize("NFKD", s)
    return "".join(c for c in s if not unicodedata.combining(c)).lower()


_LINK_PORTAL = re.compile(
    r"abre_catalogo\.asp\?t=1\|(?P<site>https?://[^|]+)\|(?P<leilao>\d+)\|(?P<lote>\d+)",
    re.IGNORECASE,
)


def link_direto_do_lote(link_portal: str) -> str:
    """
    O portal usa links como  abre_catalogo.asp?t=1|http://site|65404|32697493
    (com caracteres '|'), que o servidor rejeita quando o navegador os codifica
    ("The custom error module does not recognize this error").
    Convertemos para a página do lote no site da própria casa de leilões:
        http://site/peca.asp?ID=32697493
    """
    m = _LINK_PORTAL.search(link_portal)
    if not m:
        return link_portal.replace("|", "%7C")
    return f"{m.group('site').rstrip('/')}/peca.asp?ID={m.group('lote')}"


def md_seguro(texto: str) -> str:
    """Escapa caracteres que quebrariam o texto de um link em Markdown."""
    return re.sub(r"([\[\]])", r"\\\1", texto)


PALAVRAS_SELO = re.compile(
    r"\b(selos?|filatel\w*|rhm|yvert|scott|michel|stanley gibbons|sg\s?\d+|"
    r"bloco|blocos|s[ée]rie|olho[s]? de boi|olho de cabra|olho de gato|"
    r"stamps?|briefmarken?|timbres?|sellos?|francobolli|"
    r"carimbo|obliterad\w*|charneira|picote|dentado|sobrecarga|"
    r"inteiro postal|primeiro dia|fdc|aerograma|vinheta)\b",
    re.IGNORECASE,
)


# Itens que aparecem em categorias/lojas de filatelia mas não são selos.
# OUTROS: descartados, a menos que o título cite selo/filatelia (ex.: "moeda e selo")
NAO_SELO = re.compile(
    r"\b(moedas?|c[eé]dulas?|notas? de|medalhas?|condecora\w*|rel[oó]gios?|joias?|semi ?joias?|"
    r"an[eé]is|anel|colar|brincos?|pulseiras?|prata|ouro|pedras?|esmeralda|turmalina|ametista|"
    r"bonecos?|brinquedos?|figuras?|a[cç][aã]o|disco|vinil|lp|gibis?|hq|quadrinhos|revistas?|"
    r"livros?|quadros?|pinturas?|gravuras?|esculturas?|tapetes?|m[oó]veis|mesa|cadeira|"
    r"porcelana|vidro|cristal|canetas?|fichas?|tokens?|kit misto)\b",
    re.IGNORECASE,
)
# ACESSÓRIOS: sempre descartados (mesmo citando "selos", ex.: "álbum para selos")
ACESSORIOS = re.compile(
    r"\b([aá]lbum|[aá]lbuns|classificador\w*|pin[cç]as?|lupas?|protetor\w*|suplementos?|"
    r"cat[aá]logos?|estojos?|acess[oó]rios?|material filat[eé]lico|charneiras?\s+(para|de)|"
    r"folhas? (para|de) [aá]lbum)\b",
    re.IGNORECASE,
)


def passa_filtro_selo(titulo: str, rigor: str) -> bool:
    if rigor.startswith("Desligado"):
        return True
    t = norm(titulo)
    if ACESSORIOS.search(t):
        return False
    tem_selo = bool(PALAVRAS_SELO.search(t))
    if rigor.startswith("Estrito"):
        return tem_selo
    # Moderado: descarta o que parece outro tipo de item, a menos que cite selo/filatelia
    return tem_selo or not NAO_SELO.search(t)


@st.cache_data(ttl=600, show_spinner=False)
def philasearch_leiloes_atuais():
    """
    Lê a home do Philasearch e devolve {id_da_casa: {números dos leilões atuais}}
    a partir da seção "Current Auctions" (leilões marcados como atuais/ao vivo).
    Um número None significa "leilão atual sem número identificável no título".
    """
    try:
        r = requests.get("https://www.philasearch.com/en", headers=HEADERS, timeout=TIMEOUT)
        r.raise_for_status()
    except requests.RequestException:
        return None
    h = r.text
    ini = h.find("Current Auctions")
    fins = [x for x in (h.find("Special fixed priced sales"),
                        h.find("Auction Results and After Auction Sale")) if x > ini]
    if ini < 0 or not fins:
        return None
    trecho = h[ini:min(fins)]
    atuais: dict[str, set] = {}
    for m in re.finditer(r'href="[^"]*?/cat/\d+_(\d+)/[^"]*"[^>]*>(.*?)</a>', trecho, re.S):
        casa, txt = m.group(1), re.sub(r"<[^>]+>", " ", m.group(2))
        # os títulos começam com a data ("Oct. 4th, 2026 ..."): ignora tudo até o ano
        depois_do_ano = re.split(r"20\d\d", txt, maxsplit=1)[-1]
        n = re.search(r"(\d+)\s*(?:st|nd|rd|th)\b|#(\d+)", depois_do_ano)
        num = int(n.group(1) or n.group(2)) if n else None
        atuais.setdefault(casa, set()).add(num)
    return atuais or None


def philasearch_ativo(link: str, atuais) -> bool:
    """Mantém o lote só se a casa tem leilão atual (e, quando dá para comparar, o mesmo nº)."""
    if atuais is None:
        return True
    m = re.search(r"/i_(\d+)-A(\d+)-", link)
    if not m:
        return True  # formato desconhecido: na dúvida, mantém
    casa, num = m.group(1), int(m.group(2))
    nums = atuais.get(casa)
    if nums is None:
        return False          # casa sem leilão em andamento
    if None in nums or num >= 100000:
        return True           # não dá para comparar com segurança
    return num in nums


def termos_no_titulo(titulo: str, consulta: str) -> bool:
    t = norm(titulo)
    return all(term in t for term in norm(consulta).split())


@st.cache_resource
def _robots(base: str):
    """
    Lê o robots.txt com o mesmo User-Agent do app. Se não houver um robots.txt acessível
    (erro de rede, 403, 404...), considera que não há restrições declaradas — o
    RobotFileParser padrão trata 401/403 como "proibido tudo", o que bloqueava sites
    que apenas recusam o User-Agent do Python.
    """
    try:
        r = requests.get(urljoin(base, "/robots.txt"), headers=HEADERS, timeout=TIMEOUT)
    except requests.RequestException:
        return None
    if r.status_code != 200:
        return None
    rp = RobotFileParser()
    rp.parse(r.text.splitlines())
    return rp


def permitido(url: str) -> bool:
    p = urlparse(url)
    rp = _robots(f"{p.scheme}://{p.netloc}")
    return True if rp is None else rp.can_fetch(HEADERS["User-Agent"], url)


def montar_url(site: Site, consulta: str, pag) -> str:
    """
    Monta a URL de busca. Marcadores aceitos no modelo:
      {q}   termo com '+' no lugar de espaço (ex.: penny+black)
      {qp}  termo com '%20' no lugar de espaço (para termos dentro do caminho da URL)
      {pag} número da página
    """
    return site.url_busca.format(
        q=quote_plus(consulta, encoding=site.encoding, errors="replace"),
        qp=quote(consulta, encoding=site.encoding, errors="replace"),
        pag=pag,
    )


@st.cache_resource
def _sessao(base_home: str) -> requests.Session:
    """Sessão com cookies: alguns sites redirecionam para a home quem chega sem cookie."""
    sess = requests.Session()
    sess.headers.update(HEADERS)
    try:
        sess.get(base_home, timeout=TIMEOUT)
    except requests.RequestException:
        pass
    return sess


def baixar_com_navegador(url: str) -> tuple[str, str]:
    """Abre a página num Chromium invisível (Playwright) e devolve (html, url_final)."""
    from playwright.sync_api import sync_playwright  # pip install playwright

    with sync_playwright() as pw:
        navegador = pw.chromium.launch(headless=True)
        try:
            pagina = navegador.new_page(user_agent=HEADERS["User-Agent"], locale="en-US")
            pagina.goto(url, wait_until="domcontentloaded", timeout=30000)
            try:
                pagina.wait_for_load_state("networkidle", timeout=8000)
            except Exception:
                pass  # alguns sites nunca ficam "ociosos"; segue com o que carregou
            return pagina.content(), pagina.url
        finally:
            navegador.close()


def buscar_pagina(site: Site, consulta: str, pag: int, respeitar_robots: bool,
                  usar_navegador: bool = False):
    """Baixa uma página de resultados e devolve (itens, diagnóstico)."""
    url = montar_url(site, consulta, "" if (site.pag_1_vazia and pag == 1) else pag)
    if pag >= 2 and site.param_pagina:
        url += site.param_pagina.format(pag=pag)
    diag = {"url": url, "url_final": "", "status": None, "bytes": 0, "links": 0,
            "erro": None, "html": ""}

    if respeitar_robots and not permitido(url):
        diag["erro"] = "O robots.txt do site não permite leitura automática desta página (opção \"Respeitar robots.txt\" em Avançado)"
        return [], diag
    try:
        if site.js and usar_navegador:
            texto_html, diag["url_final"] = baixar_com_navegador(url)
            conteudo = texto_html.encode("utf-8")
            diag["status"] = 200
            diag["bytes"] = len(conteudo)
        else:
            if site.sessao:
                p_ = urlparse(url)
                home = f"{p_.scheme}://{p_.netloc}/en"
                r = _sessao(home).get(url, headers={**HEADERS, "Referer": home}, timeout=TIMEOUT)
            else:
                r = requests.get(url, headers=HEADERS, timeout=TIMEOUT)
            diag["url_final"] = r.url
            diag["status"] = r.status_code
            diag["bytes"] = len(r.content)
            r.raise_for_status()
            conteudo, texto_html = r.content, r.text
    except ImportError:
        diag["erro"] = ("Playwright não instalado. Rode: pip install playwright && "
                        "playwright install chromium")
        return [], diag
    except requests.RequestException as e:
        diag["erro"] = f"Falha na requisição: {e}"
        return [], diag
    except Exception as e:  # falhas do navegador automatizado
        diag["erro"] = f"Falha no navegador automatizado: {e}"
        return [], diag

    # o portal declara ISO-8859-1; deixa o requests/bs4 detectarem
    soup = BeautifulSoup(conteudo, "html.parser")
    diag["html"] = texto_html[:3000]

    itens, vistos = [], set()
    for a in soup.select(site.seletor_link):
        href = a.get("href", "")
        if not href or href.startswith(("javascript", "#")):
            continue
        link = urljoin(url, href)
        if site.href_exige and not re.search(site.href_exige, link, re.IGNORECASE):
            continue
        if site.href_exclui and re.search(site.href_exclui, link, re.IGNORECASE):
            continue
        img = a.find("img")
        titulo = (a.get("title") or a.get_text(" ", strip=True)
                  or (img.get("alt") if img else "") or "").strip()
        titulo = re.split(r"\s::\s", titulo)[0]  # "título :: vendedor"
        titulo = re.sub(r"(Opens in a new window or tab|New Listing|Novo anúncio)", "",
                        titulo, flags=re.IGNORECASE).strip()
        if len(titulo) < 6 or link in vistos:  # ignora "Ver", "Comprar", "Home"...
            continue
        vistos.add(link)
        itens.append(
            {
                "Título": re.sub(r"\s+", " ", titulo),
                "Link": link_direto_do_lote(link),
                "Link do portal": link,
            }
        )
    # Fallback: o seletor CSS não achou nada (HTML mudou?). Varre todos os links da página
    # e mantém só os que citam todos os termos buscados no texto/título.
    if not itens:
        for a in soup.find_all("a", href=True):
            href = a["href"]
            if href.startswith(("javascript", "#", "mailto:", "tel:")):
                continue
            img = a.find("img")
            titulo = (a.get("title") or a.get_text(" ", strip=True)
                      or (img.get("alt") if img else "") or "").strip()
            link = urljoin(url, href)
            if len(titulo) < 6 or link in vistos or not termos_no_titulo(titulo, consulta):
                continue
            vistos.add(link)
            itens.append({"Título": re.sub(r"\s+", " ", titulo),
                          "Link": link_direto_do_lote(link), "Link do portal": link})
        diag["fallback"] = bool(itens)
    diag["links"] = len(itens)
    return itens, diag


ARQUIVO_EXTRAS = Path(__file__).with_name("sites_extras.json")


def site_a_partir_de_url(nome: str, url: str, termo: str, js: bool = False) -> tuple[Site, str]:
    """
    Cria um Site a partir de uma URL de resultados de busca copiada do navegador.
    Troca o termo usado no exemplo por {q} e o número da página por {pag}.
    Devolve (site, explicação).
    """
    url = url.strip()
    termo = termo.strip()
    u = url.replace("{", "{{").replace("}", "}}")  # protege chaves para o .format()
    encoding, achou = "utf-8", False
    variantes = [  # (texto na URL, codificação, marcador)
        (quote_plus(termo), "utf-8", "{q}"), (quote(termo), "utf-8", "{qp}"),
        (quote_plus(termo, encoding="latin-1"), "latin-1", "{q}"),
        (quote(termo, encoding="latin-1"), "latin-1", "{qp}"),
        (termo, "utf-8", "{q}"), (termo.replace(" ", "-"), "utf-8", "{q}"),
        (termo.replace(" ", "_"), "utf-8", "{q}"),
    ]
    for v, enc, ph in sorted(variantes, key=lambda x: -len(x[0])):
        if v and re.search(re.escape(v), u, re.IGNORECASE):
            u = re.sub(re.escape(v), ph, u, count=1, flags=re.IGNORECASE)
            encoding, achou = enc, True
            break
    m = re.search(r"([?&](?:page|pag|pagina|pg|p|_pgn|mpage|pageno|pn)=)(\d*)", u, re.IGNORECASE)
    usa_pag = bool(m)
    if m:
        u = u[:m.start(2)] + "{pag}" + u[m.end(2):]
    site = Site(
        nome=nome, grupo=nome, url_busca=u, seletor_link="a[href]", encoding=encoding,
        usa_paginas=usa_pag, js=js, busca_local=True, extra=True,
    )
    msg = (
        f"Modelo criado: `{u}`  \n"
        + ("✔ termo de busca reconhecido → {q}. " if achou else
           "⚠ não encontrei o termo na URL; confira se digitou o mesmo termo usado nela. ")
        + ("✔ paginação reconhecida → {pag}." if usa_pag else "Sem paginação reconhecida (só a 1ª página).")
    )
    return site, msg


def carregar_extras() -> list[Site]:
    try:
        dados = json.loads(ARQUIVO_EXTRAS.read_text(encoding="utf-8"))
        campos = set(Site.__dataclass_fields__)
        out = []
        for d in dados:
            try:
                out.append(Site(**{k: v for k, v in d.items() if k in campos}))
            except Exception:
                continue          # um registro ruim não derruba os demais
        return out
    except Exception:
        return []


def salvar_extras(sites: list[Site]) -> None:
    try:
        ARQUIVO_EXTRAS.write_text(
            json.dumps([asdict(x) for x in sites if x.extra], ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
    except OSError:
        pass


def url_busca_manual(site: Site, consulta: str) -> str:
    return montar_url(site, consulta, "" if site.pag_1_vazia else 1)


def lote_ativo(link: str) -> bool:
    """Confere na página do lote se ele não aparece como vendido/encerrado."""
    try:
        r = requests.get(link, headers=HEADERS, timeout=TIMEOUT)
        return not ENCERRADO_NA_PAGINA.search(r.text)
    except requests.RequestException:
        return True  # na dúvida, mantém


def novo_stats(nome: str) -> dict:
    return {
        "Site": nome, "HTTP": None, "Bytes": 0, "Links achados": 0,
        "Descartados (título)": 0, "Descartados (não é selo)": 0,
        "Descartados (encerrado)": 0, "Mantidos": 0, "Tempo (s)": 0.0,
        "Erro": "", "Nota": "", "URL": "", "URL final": "", "_html": "", "_amostra": [],
    }


def buscar_sitemap(site: Site, consulta: str, respeitar_robots: bool,
                   t0: float, limite_s: int):
    """
    Lê o sitemap.xml público do site (a lista de todas as páginas de produto) e devolve os
    produtos cujo endereço (slug) contém todos os termos buscados. É a forma mais confiável
    de achar um produto numa loja grande sem percorrer o catálogo inteiro.
    Devolve (itens, erro). itens=None significa que o sitemap não pôde ser usado.
    """
    if respeitar_robots and not permitido(site.sitemap):
        return None, "o robots.txt não permite ler o sitemap"
    tokens = norm(consulta).split()
    urls, fila, lidos = [], [site.sitemap], 0
    t_sitemap = time.monotonic()
    orcamento = min(20, limite_s * 0.5)   # o sitemap pode ocupar no máximo metade do tempo
    while fila and lidos < 12 and time.monotonic() - t_sitemap < orcamento:
        u = fila.pop(0)
        lidos += 1
        try:
            r = requests.get(u, headers=HEADERS, timeout=TIMEOUT)
            r.raise_for_status()
        except requests.RequestException as e:
            if lidos == 1:
                return None, f"sitemap indisponível: {e}"
            continue
        for loc in re.findall(r"<loc>\s*(?:<!\[CDATA\[)?\s*(.*?)\s*(?:\]\]>)?\s*</loc>", r.text, re.S):
            loc = html.unescape(loc.strip())
            caminho = urlparse(loc).path.lower()
            if caminho.endswith(".gz"):
                continue
            if caminho.endswith(".xml") or "sitemap" in caminho:
                fila.append(loc)          # sitemap dentro de sitemap
            else:
                urls.append(loc)
    if not urls:
        return None, "sitemap sem endereços de páginas"

    candidatos = []
    for loc in urls:
        if site.href_exige and not re.search(site.href_exige, loc, re.IGNORECASE):
            continue
        slug = urlparse(loc).path.rstrip("/").rsplit("/", 1)[-1]
        slug_n = norm(slug).replace("-", " ").replace("_", " ")
        if all(t in slug_n for t in tokens):
            candidatos.append(loc)
    candidatos = list(dict.fromkeys(candidatos))[:40]

    def titulo_de(loc: str) -> str:
        slug = urlparse(loc).path.rstrip("/").rsplit("/", 1)[-1].replace("-", " ").strip()
        if time.monotonic() - t0 > limite_s:
            return slug.capitalize()
        try:
            r = requests.get(loc, headers=HEADERS, timeout=(5, 8))
            sp = BeautifulSoup(r.text, "html.parser")
            og = sp.find("meta", attrs={"property": "og:title"})
            t = (og.get("content") if og else "") or (sp.title.get_text() if sp.title else "")
            t = re.sub(r"\s+", " ", t).strip()
            return t or slug.capitalize()
        except requests.RequestException:
            return slug.capitalize()

    with ThreadPoolExecutor(max_workers=6) as ex:
        titulos = list(ex.map(titulo_de, candidatos))
    itens = [{"Título": t, "Link": loc, "Link do portal": loc}
             for loc, t in zip(candidatos, titulos)]
    return itens, None


def processar_site(site: Site, consulta: str, max_paginas: int,
                   respeitar_robots: bool, exigir_titulo: bool, verificar: bool,
                   rigor: str = "Moderado", usar_navegador: bool = False,
                   limite_s: int = 45):
    t0 = time.monotonic()
    resultados, vistos = [], set()
    stats = novo_stats(site.nome)
    usando_reserva = False
    if site.url_reserva and respeitar_robots:
        url_principal = montar_url(site, consulta, "")
        if not permitido(url_principal):
            # A busca do site não é liberada para robôs: lê o sitemap e, se preciso, o catálogo.
            usando_reserva = True
            stats["Nota"] = ("robots.txt barra a busca do site; usei o sitemap/catálogo e "
                             "filtrei o termo aqui")
            site = replace(site, url_busca=site.url_reserva, usa_paginas=True,
                           param_pagina="", pag_1_vazia=False)
    paginas = max(max_paginas, site.min_paginas) if site.usa_paginas else 1
    atuais = philasearch_leiloes_atuais() if site.filtro_ativos == "philasearch" else None

    def filtrar(novos):
        """Aplica os filtros e acumula os lotes aprovados em 'resultados'."""
        for i in novos:
            if (exigir_titulo or site.busca_local) and not termos_no_titulo(i["Título"], consulta):
                stats["Descartados (título)"] += 1
                if len(stats["_amostra"]) < 6:
                    stats["_amostra"].append(i["Título"][:140])
                continue
            if site.filtro_ativos == "philasearch" and not philasearch_ativo(i["Link"], atuais):
                stats["Descartados (encerrado)"] += 1
                continue
            if not passa_filtro_selo(i["Título"], rigor):
                stats["Descartados (não é selo)"] += 1
                continue
            if ENCERRADO.search(i["Título"]):
                stats["Descartados (encerrado)"] += 1
                continue
            resultados.append({"Site": site.grupo or site.nome, **i})

    # 1) sitemap público (se o site tiver um configurado)
    usou_sitemap = False
    if site.sitemap and (usando_reserva or not site.url_reserva):
        itens_sm, erro_sm = buscar_sitemap(site, consulta, respeitar_robots, t0, limite_s)
        stats["URL"] = site.sitemap
        if itens_sm is not None:
            usou_sitemap = True
            stats["HTTP"] = 200
            stats["Links achados"] = len(itens_sm)
            for i in itens_sm:
                vistos.add(i["Link"])
            filtrar(itens_sm)
        else:
            stats["Erro"] = f"Sitemap não usado ({erro_sm}); lendo o catálogo."

    # 2) catálogo/resultados paginados
    for pag in range(1, (0 if usou_sitemap else paginas) + 1):
        if time.monotonic() - t0 > limite_s:
            stats["Erro"] = (f"Tempo máximo de {limite_s}s atingido antes da página {pag} "
                             "(resultados parciais)")
            break
        itens, d = buscar_pagina(site, consulta, pag, respeitar_robots, usar_navegador)
        if pag == 1:
            stats.update({"HTTP": d["status"], "Bytes": d["bytes"],
                          "URL": d["url"], "URL final": d["url_final"],
                          "_html": d["html"]})
            if not d["erro"] and site.sitemap:
                stats["Erro"] = ""        # o catálogo funcionou: apaga o aviso do sitemap
        if d["erro"]:
            if pag == 1:
                stats["Erro"] = d["erro"]
            else:  # páginas seguintes são opcionais: guarda só uma nota
                stats["Nota"] = f"Páginas seguintes não lidas (página {pag}: {d['erro'][:70]})"
            break
        novos = [i for i in itens if i["Link"] not in vistos]
        if not novos:  # página repetida/vazia -> fim da paginação
            break
        for i in novos:
            vistos.add(i["Link"])
        stats["Links achados"] += len(novos)
        filtrar(novos)
        time.sleep(0.5)

    if verificar and resultados and time.monotonic() - t0 < limite_s:
        # confere no máximo 40 lotes por site, para não travar
        with ThreadPoolExecutor(max_workers=6) as ex:
            ok = list(ex.map(lambda x: lote_ativo(x["Link"]), resultados[:40]))
        ok += [True] * (len(resultados) - len(ok))
        antes = len(resultados)
        resultados = [r for r, o in zip(resultados, ok) if o]
        stats["Descartados (encerrado)"] += antes - len(resultados)

    stats["Mantidos"] = len(resultados)
    stats["Tempo (s)"] = round(time.monotonic() - t0, 1)
    return resultados, stats


# ----------------------------------------------------------------------------
# Interface
# ----------------------------------------------------------------------------
# ----------------------------------------------------------------------------
# Exportação dos resultados: Excel (.xlsx) e Word (.docx)
#   pip install openpyxl python-docx
# ----------------------------------------------------------------------------
_CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _limpa(texto) -> str:
    """Remove caracteres de controle, que o XML de Excel/Word não aceita."""
    return _CTRL.sub("", str(texto if texto is not None else ""))


def links_do_lote(row):
    """(link principal, link secundário ou None, rótulo do secundário) de um lote."""
    direto, portal = row["Link"], row["Link do portal"]
    if direto == portal:                       # site comum: um link só
        return direto, None, ""
    try:
        direto_falhou = row.get("ok_direto") in (False,)
    except TypeError:
        direto_falhou = False
    if direto_falhou:                          # a página direta deu erro no teste
        return portal, direto, "link direto"
    return direto, portal, "via LeilõesBR"      # padrão: página do lote na casa de leilões


def _resumo_sites(stats_list) -> list[tuple[str, int, str]]:
    return [(x["Site"], x["Mantidos"], x["Erro"]) for x in stats_list]


def gerar_excel(df, consulta: str, stats_list) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    navy, laranja, azul = "0B1F3A", "D9480F", "1E4D8C"
    fonte = lambda **kw: Font(name="Arial", size=kw.pop("size", 10), **kw)
    fino = Side(style="thin", color="DCE4EF")

    wb = Workbook()
    ws = wb.active
    ws.title = "Resultados"
    cab = ["Site", "Título", "Abrir lote", "Link principal", "Link alternativo"]
    ws.append(cab)
    for c in ws[1]:
        c.font = fonte(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=navy)
        c.alignment = Alignment(vertical="center", horizontal="left")
    ws.row_dimensions[1].height = 22

    for i, (_, row) in enumerate(df.iterrows(), start=2):
        principal, secundario, _rot = links_do_lote(row)
        ws.append([_limpa(row["Site"]), _limpa(row["Título"]), "Abrir lote",
                   _limpa(principal), _limpa(secundario or "")])
        t = ws.cell(row=i, column=2)
        t.data_type = "s"            # título que comece com "=" nunca vira fórmula
        c = ws.cell(row=i, column=3)
        c.hyperlink = _limpa(principal)
        c.font = fonte(color=azul, underline="single")
        for col in (1, 2, 4, 5):
            ws.cell(row=i, column=col).font = fonte()
        for col in range(1, 6):
            cell = ws.cell(row=i, column=col)
            cell.alignment = Alignment(vertical="top", wrap_text=col in (2,))
            cell.border = Border(bottom=fino)

    for col, w in zip("ABCDE", (34, 78, 14, 58, 48)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:E{max(ws.max_row, 2)}"

    r = wb.create_sheet("Busca")
    r["A1"], r["B1"] = "eFilatelia", "Resumo da busca"
    r["A1"].font = fonte(bold=True, size=14, color=laranja)
    r["B1"].font = fonte(bold=True, size=14, color=navy)
    dados = [
        ("Termo buscado", _limpa(consulta)),
        ("Data e hora", datetime.now().strftime("%d/%m/%Y %H:%M")),
        ("Lotes encontrados", len(df)),
    ]
    for k, (rot, val) in enumerate(dados, start=3):
        r.cell(row=k, column=1, value=rot).font = fonte(bold=True, color=navy)
        r.cell(row=k, column=2, value=val).font = fonte()
        r.cell(row=k, column=2).alignment = Alignment(horizontal="left")
    base = len(dados) + 4
    for j, h in enumerate(("Site", "Lotes", "Observação"), start=1):
        c = r.cell(row=base, column=j, value=h)
        c.font = fonte(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=navy)
    for k, (site, n, erro) in enumerate(_resumo_sites(stats_list), start=base + 1):
        r.cell(row=k, column=1, value=_limpa(site)).font = fonte()
        r.cell(row=k, column=2, value=n).font = fonte()
        r.cell(row=k, column=2).alignment = Alignment(horizontal="left")
        ob = r.cell(row=k, column=3, value=_limpa(erro) or "—")
        ob.font = fonte()
        ob.alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABC", (42, 28, 70)):
        r.column_dimensions[col].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _docx_hyperlink(par, texto: str, url: str, negrito=False, cor="1E4D8C", tam=10):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    r_id = par.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    for tag, attrs in (("w:rFonts", {"w:ascii": "Arial", "w:hAnsi": "Arial"}),
                       ("w:b", {}) if negrito else (None, None),
                       ("w:color", {"w:val": cor}),
                       ("w:sz", {"w:val": str(int(tam * 2))}),
                       ("w:u", {"w:val": "single"})):
        if tag is None:
            continue
        el = OxmlElement(tag)
        for k, v in attrs.items():
            el.set(qn(k), v)
        rpr.append(el)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = texto
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    hl.append(run)
    par._p.append(hl)


def gerar_word(df, consulta: str, stats_list) -> bytes:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    navy, laranja = RGBColor(0x0B, 0x1F, 0x3A), RGBColor(0xD9, 0x48, 0x0F)

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)          # A4
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(2)

    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    for nome, tam, cor in (("Heading 1", 18, navy), ("Heading 2", 13, navy)):
        st_ = doc.styles[nome]
        st_.font.name, st_.font.size, st_.font.bold, st_.font.color.rgb = "Arial", Pt(tam), True, cor
        rf = st_.element.rPr.rFonts   # tira a fonte de tema, que sobrepõe o Arial
        for att in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rf.get(qn(att)) is not None:
                del rf.attrib[qn(att)]

    def borda(par, lado="bottom", cor="FFB37A", estilo="dotted", sz="18"):
        ppr = par._p.get_or_add_pPr()
        bd = OxmlElement("w:pBdr")
        b = OxmlElement(f"w:{lado}")
        for k, v in (("w:val", estilo), ("w:sz", sz), ("w:space", "4"), ("w:color", cor)):
            b.set(qn(k), v)
        bd.append(b)
        ppr.append(bd)

    # marca
    p = doc.add_paragraph()
    r1 = p.add_run("e")
    r1.font.size, r1.font.bold, r1.font.color.rgb = Pt(24), True, laranja
    r2 = p.add_run("Filatelia")
    r2.font.size, r2.font.bold, r2.font.color.rgb = Pt(24), True, navy
    borda(p)

    doc.add_heading("Resultados da busca", level=1)
    sub = doc.add_paragraph()
    a = sub.add_run("Termo buscado: ")
    a.font.color.rgb = RGBColor(0x5A, 0x6F, 0x92)
    b_ = sub.add_run(f"“{_limpa(consulta)}”")
    b_.bold = True
    c_ = sub.add_run(f"   ·   {datetime.now().strftime('%d/%m/%Y %H:%M')}")
    c_.font.color.rgb = RGBColor(0x5A, 0x6F, 0x92)

    n = len(df)
    n_sites = df["Site"].nunique()
    res = doc.add_paragraph()
    num = res.add_run(str(n))
    num.font.size, num.font.bold, num.font.color.rgb = Pt(20), True, laranja
    res.add_run(f"  {'lote' if n == 1 else 'lotes'} em leilões em andamento, "
                f"em {n_sites} {'site' if n_sites == 1 else 'sites'}.")

    def bordas_tabela(tabela, cor="DCE4EF"):
        tblpr = tabela._tbl.tblPr
        bd = OxmlElement("w:tblBorders")
        for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{lado}")
            for k, v in (("w:val", "single"), ("w:sz", "4"), ("w:space", "0"), ("w:color", cor)):
                el.set(qn(k), v)
            bd.append(el)
        tblpr.append(bd)

    def sombra(celula, cor):
        tcpr = celula._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        for k, v in (("w:val", "clear"), ("w:color", "auto"), ("w:fill", cor)):
            shd.set(qn(k), v)
        tcpr.append(shd)

    larg = (Cm(13.6), Cm(3.4))
    for nome_site, g in df.groupby("Site", sort=False):
        doc.add_heading(f"{_limpa(nome_site)} ({len(g)})", level=2)
        tab = doc.add_table(rows=1, cols=2)
        tab.alignment = WD_TABLE_ALIGNMENT.CENTER
        tab.autofit = False
        for col_, w_ in zip(tab.columns, larg):   # larguras na grade (Word e LibreOffice)
            col_.width = w_
        bordas_tabela(tab)
        for cel, txt, w in zip(tab.rows[0].cells, ("Lote", "Outro link"), larg):
            cel.width = w
            sombra(cel, "0B1F3A")
            run = cel.paragraphs[0].add_run(txt)
            run.bold, run.font.color.rgb, run.font.size = True, RGBColor(255, 255, 255), Pt(9.5)
        for _, row in g.iterrows():
            principal, secundario, rotulo = links_do_lote(row)
            cels = tab.add_row().cells
            cels[0].width, cels[1].width = larg
            _docx_hyperlink(cels[0].paragraphs[0], _limpa(row["Título"]), _limpa(principal))
            if secundario:
                _docx_hyperlink(cels[1].paragraphs[0], rotulo, _limpa(secundario),
                                cor="C2410C", tam=9)
        doc.add_paragraph()

    problemas = [(s_, e_) for s_, _, e_ in _resumo_sites(stats_list) if e_]
    if problemas:
        doc.add_heading("Sites que não puderam ser lidos", level=2)
        for s_, e_ in problemas:
            li = doc.add_paragraph(style="List Bullet")
            li.add_run(f"{_limpa(s_)}: ").bold = True
            li.add_run(_limpa(e_))

    fim = doc.add_paragraph()
    borda(fim, lado="top")
    f = fim.add_run("Gerado pelo eFilatelia. Lances e compras acontecem no site do leiloeiro.")
    f.font.size, f.font.color.rgb = Pt(8.5), RGBColor(0x5A, 0x6F, 0x92)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _slug(texto: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", norm(texto)).strip("-")[:40] or "busca"


# ----------------------------------------------------------------------------
# Identidade visual eFilatelia: branco, laranja e azul-marinho
# ----------------------------------------------------------------------------
LOGO_SVG = r"""<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 296 66" role="img" aria-label="eFilatelia">
  <title>eFilatelia</title>
  <g transform="translate(4,4)">
  <rect width="48" height="58" fill="#0B1F3A"/>
  <rect x="6" y="6" width="36" height="46" rx="1.5" fill="#FFFFFF"/>
  <g fill="#FFFFFF"><circle cx="4" cy="0" r="2.5"/><circle cx="4" cy="58" r="2.5"/><circle cx="12" cy="0" r="2.5"/><circle cx="12" cy="58" r="2.5"/><circle cx="20" cy="0" r="2.5"/><circle cx="20" cy="58" r="2.5"/><circle cx="28" cy="0" r="2.5"/><circle cx="28" cy="58" r="2.5"/><circle cx="36" cy="0" r="2.5"/><circle cx="36" cy="58" r="2.5"/><circle cx="44" cy="0" r="2.5"/><circle cx="44" cy="58" r="2.5"/><circle cx="0" cy="5" r="2.5"/><circle cx="48" cy="5" r="2.5"/><circle cx="0" cy="13" r="2.5"/><circle cx="48" cy="13" r="2.5"/><circle cx="0" cy="21" r="2.5"/><circle cx="48" cy="21" r="2.5"/><circle cx="0" cy="29" r="2.5"/><circle cx="48" cy="29" r="2.5"/><circle cx="0" cy="37" r="2.5"/><circle cx="48" cy="37" r="2.5"/><circle cx="0" cy="45" r="2.5"/><circle cx="48" cy="45" r="2.5"/><circle cx="0" cy="53" r="2.5"/><circle cx="48" cy="53" r="2.5"/></g>
  <path d="M13 29H35A11 11 0 1 0 32.4 36.1" fill="none" stroke="#F26B1D" stroke-width="5.2" stroke-linecap="round" stroke-linejoin="round"/>
  <circle cx="40" cy="13" r="8.5" fill="none" stroke="#FFB37A" stroke-width="1.7"/>
  <path d="M51 6q2.75 -2.2 5.5 0t5.5 0t5.5 0t5.5 0" fill="none" stroke="#F26B1D" stroke-width="1.7" stroke-linecap="round"/><path d="M51 11q2.75 -2.2 5.5 0t5.5 0t5.5 0t5.5 0" fill="none" stroke="#F26B1D" stroke-width="1.7" stroke-linecap="round"/><path d="M51 16q2.75 -2.2 5.5 0t5.5 0t5.5 0t5.5 0" fill="none" stroke="#F26B1D" stroke-width="1.7" stroke-linecap="round"/><path d="M51 21q2.75 -2.2 5.5 0t5.5 0t5.5 0t5.5 0" fill="none" stroke="#F26B1D" stroke-width="1.7" stroke-linecap="round"/>
  </g>
  <text x="90" y="42" class="wm"><tspan fill="#F26B1D">e</tspan><tspan fill="#0B1F3A">Filatelia</tspan></text>
  <text x="92" y="57" class="tg" fill="#1E4D8C">leilões de selos em uma só busca</text>
</svg>"""
ICONE_PNG_B64 = "iVBORw0KGgoAAAANSUhEUgAAAEAAAABACAYAAACqaXHeAAATaUlEQVR42u1beXRV1dX/7XPvffdNeUmYkgAhCIqIOGAAEayVakEsUhmcB7RSLWIUrK18aKFI7SrSooDiUG0RxYkKKIiWmVKFD0EmGSMkJBAJGIbkTfe+e8/+/rj3ZYLAi4Sub6161jrrDeeec8/ZZ589/M7ewA/lh/JfXej7dmRmAgAi4lTbkv+frq2pxjunhZnFqSbRmOdq/27o+/d977lePrkTaMHMbZIbMGHChNqTa8vMzWo9n+zjZeZ2zKwCwIQJEwS4erxsZm7dwHjtmDlUi2uTfYLM3K7+3M7lzivu5yBm/o6Zo7ZtTwEARREQRGDm15g5ysyHmfnGWn3zmHmr27aGmTOFoGTbA8xcycwRZh6XHI+ZNWZe5PYpZuYetca7mJkL3bZFzOwh5/3njggrV65UH332r1lxw1zJNcXesHn7lUDztF3fFF8npaxuSCQSqx9+4rlsZOanR+PGdGZmy7JsZubyI0dHIf2yjCkvvp1nJhJf1xrP+OjTNRcD54eKS8uGMTPbtu00GMbcq28paNm2S79m8bj5JjOz5baVlpUPyfvxcG/tY9KkwpKZRZdetzTztetT/q+1X9nMbDOzNM2EvLL/vXGkXXziJ0MeMphZJuvyNetlsP3VETS7rHLyi7MSyT7MbD8ydnICGZdWZl/80+jOPfuS/8ujx0/IC3r+PIaMyyqH3veEWbvPPxYulZ7WPauUrCuq3vpgkVWrzbhlxDgTwa7jmZny8x/UUl2YkioBJk6cyNNn/CVr1b/W/2btV1+rXTt3RDgSpT++8AYtWrJG9YWC+r7iA0okGqPcNlm0q7CInnxmOpWWHdJ0r65v2/mNaJvTirxeneZ+vJReeG2OIEXVT1RFtOKSMup0fjv67uhxmvCnl2nNus2qP+jX9+wtVhRFodZZLeirrTvp8d9Pp8OHj3uYhGfLjm9Ep/PakqoKenXWh/Jv7y7Uuve4ZMeIu25e/O23G2WTqUFmJiJiZv4lIH9bXFreYczTU2jBwuWUlp6GcCSKQMDPQggwM6LRGKUF/DBME7Zk+Lw6M4BEIkGWZSPg9yIcicHj0VhTVTCAaCxGuuYBESFuGPD7vQwQpJSIRKIkVB22YSK7uQ+tmwWYAZRVhKm8IgZoKtJ8Kr8xbSJ+1u9Hht/n/SeAAiIqTc79exOg1uKzABwAoALgnXuKqGf/e2DbNjyahljcgGmaABGEEJC27XwngpTO+ykp8GwbpDiMx5IBQk0fAEIRTh9mgAQ0rxf9LsrE8N6tkN8+Da2CGoQgRGwFm0oimPTBFlzQqx/+9sLTgJQ2hFAATCGi3zKzSkTW6daoNkYJJL/YUoIZUIRAzDDQJqcVclu3gi0lAAJRkoCnoDY5I9Xflrp9GCQUkGXi3ksJI67JAYLZWLY7gnlbqhA3TCiRMgzunoM143ujyA7BjkeheP2i/lybzOhh5ods2/66cF+JvOPBsaxk5bO37ZXcoftA3razkB0Jb7OU8uyrbbMtbeYNbzEv+R1z0Vquqgrz0bDJFWGDDx2L8ANPTGaReTk/es9tHPvwcSk3vMnHjx49zszzkrZBU6pEAoA5/1jaoWOPQVGRlc/pHX4k0eJyvm7Ig7U0mOSzL+4Ye5YxfziG7bLt9Zqd9k+XrWHRohujeU++qf/ABC8dz4c++cvMxi5MNIIT6K5hPz24d/c+y+/zAq5skVJCukdCSoeFv39lMBM4dgK8ewX4wusgcroAbAPMYACWzWBmxOImIIDMTB0LN5SLKfP3IMtX1X/z5s0BZlZS3X2RysLdM6UePHR4yJ13DNIBwLJk9WEncr6edQU7n4e2g9gGdfyRKxQEQOTYwJR8H1UTDmSLZSUKM+kdLqv8/FYisoHfU5NxADNrAN5pndXynTmvPOt55smRMAwThKa2Ot3xKoqA9LaAN1RDnVNNXhAi0Ri6XdIJb7z2PMibIWV6sxeY+XqiiTIVq1CcSQC6ejQHwDAppZVIWDz4xr5ID6WBLet7eNQMsASkXau6Zye5UCMM+EIu4zUs0IkEEnED1/bJR9uclmTZHlv49BCAYalusJqi6jsKoFgI0V4Iwbv37kckGnX1OacqRJyFCyXJ76d4xiWE5gPM2BmJy8wQqoo9+0qcxSiS3OlsTVUdnpYArhFERBQ2DGOwbfPId+d/9sCfZ85WGOCUz4C03YUrgLQgDxdBHt4HjlUCigaRng3R+kJQINNZc1oWUPQFYCcARW2QEFJKBIMBLFm9HqPHTZZTB2iqOMqvAniVeYIAIM+WA5JIi9B1fTOAMUp2j7tUhQJerw4jbpx51x3zDvLoQSTWvgdr82eQR4oAIwq2TYAUkMcLCmRC6dAdWu87oOZdDuxdDZRtBXLzHc6gU3OzRyVUVUku37xSiJ9cZSCz9bNEZDNPEKmgRGpjpRRRikxf60ybK16H8dk0cEUJSPUCqgaoHpCmO1wqJbjyMBLr5yGx8WOo3QbC1+FC0O6lQMsLHGHIsg4nEABVIVSEE2jTImBPG5GvStN8W7nkjlKeAEE0UTalGpTM3D0ej097fep47/ntcxE3TIcap9v5hIH4rALE3/41EK4ABZsD3gCgaIC0AMsArISzHM0L8meAPH5YX85DdPX74MrDwIY5QPiIwwFE1ZxgsUCi0sCF2UEsGHmhaJWVyeWd772CmUfhpg1KqpagmqIzlAZgnq7rucNvH8g52c0x6J7H6ZSsUK2cGbHZY5D491ugtJau5JdgIwKwhAg2d4Qd2+DoCXC4EuQNAEIFBTJhf1eG2Fer4e1GEP+eCbTrCbTtBvJmAKqClp4oJtzeGWMHXQBvKEMkOg1FTsvzugF4Efn5AsAM190/K2coaQRlAsh1vFpL7dQhD0G/DxWR6CkEFANCwPxsGhKfvw0KtXSEIEvAjEPt1AfaVbdCOS/fabMM2GW7YW1bBut/54IjxwDdD/J4YR0pRWzLF/Df+BCo5EugcCWEqgNEuMZOoM+wLpixcAeqcq/G7/qeB8tKxFVV8wLomuqZPpMWkC4bHQIwD8AQTVN5/icrcKKyCqSqdTUNO4uXZbthfDYd5As5Ol5KgCX0wU/D02+UK9lrTSIjB2qXayF73474W4/D2r8J5AmA9ADsAztglBbBe8tE4OAWyMhRCI8HX2wpxm1jXsaBQxF0uVTBiPvvRE52Ky+AYwD+kaoaFClKPhPAbeXfHbvn7pFPJX43+WXoXg/4pPGd3+bKN8BVFYDiAcDgRAz6zU/DM+Ax5xxbpiMDkoaQbQF2AiK3K3wjZ0HkdAabUefdvjQkPp8DWb4XaHMZuFNfoH0flHnycOBwGBkt07CncK99w51j8OIbHywG0IeIlrrH1z5rAriqhADY2S2bfTBnzscGEUEVgk9y+IUCrjoKa/tykMcHEIPjYWg9h8LTb2T1M1A9gFCd70JxOEJxYDzKbA3ffS8AHp9DGNUDrqqA9dWi6mMEltDYhKIpsG0Jv98nt+/ah4IRT64lop0FBdP0VC9KRCP0H787b3Fux84d1FgsfjJvSduhUvFG8LFvnUVKCageKB17wj7wNeT+TZAlW09b7f2bQaoXSpsujpZggBQFVuE6h3tUD0ACDIJkdmC0uEEBvxfXD7sxGwBmzHjMaBIZUE8TPCilHNPzikv1p//4IuZ+vIxInEw/eWQ/2IiAtGYOUVQPjHmTXJe2ERaHogCa1zWfVfCRYiARB1RfHaPEsmxkpoeUmc+NQ7++ve4Lzn2pNYDHUsUERYqLzwLwkhCic4e8NvTUmBHw+71g267RAq5NwJGjda02Zueco5G+sZQ1SkgoYCsGjhyvo3SEEIiGI7ip/zU0ZOBPEAz4vQAGAyhIFfVuzCUCo8kBtzO90TXmEnFnqh7/aV8uIRvtn4sUnaFyAKOklDv3Fh/kP0z9K6LReF1v0BWIFGhWM/EkZwilrpF0xiodAeiayUhrAX3gEyB/uvt/jTPkDwbw8T9X84eLliMajkdddT0jid82lTNERPTXWe8uXjFp6ivbikoO+oIBH1dWhk+iuGiZ51h0LAEisGXCe+szUC7oDXKdnzO5uKQozm7bCbARBQWbQbTIcyck6jyrqgqOn6iyfzF6ktrjii6zls99uaAxV+YpO0MuJ5Qi7WIr2DzjZAdVKGDAsfAycsDHygCPF7BM2N+sh6fvLxvF+nbhWlDzdlByLjzJuapPMK9X53DUwPK5n5YDQEHBND1VTZAyJrhx40a1/MjR2+6++2adAViyBhOsowqDzaB2vR5sxgAmkDeIxPr5MJe8VIMNnGQIJZwKgI8dRPTPP0f0L0MQmdQXiQ0Lqn2LhlChSCQmLr2oA2a8Prk3M3eZMeMxo8lAUZcInm7dur3fqkXm7Ldm/kGbNHYkjPipMEHnt6fvA6C05oBtOhclug/GgmdhLn7eORonGUIaoGiQJVsRm3kfrN2fA3oAHD4Gc9krtfAAPgkTjMXiuOSi85XF77yARx64dYCE/JyZ+7nyq8kwwWwhxGDXGcLgAQ1ggkSOp5fTCfqA0eBYFSCEM3mhwFjwLKJThyCx5m3IAzvB4QrwsYOwtq9A/P1xiD4/DHbJFkfYucgvefzuuNwAJhhH36t7ICerJQwzERcQGQCGNjUmeAxACYB2mqbynn37EYnGGsAEHR3u+eko2GW7kPjXbMcdBgF6EFbhOlh7voAINHPMXem6w0bEEZ7JBRsRwBeCPmisQ0CW7nr4JEyw0MUEdY/mdZu2NYkzVEsNVgEYEosZb8x6b6Ec/dSfHalEp+xUXX13T4Xn2l84xpFlAkKAfGkgXwiciIHD34GjxwGhgIKZrvMEcPgYKNQCvodnQzm/Z4OQmJQSgWAAS1avk4+MnczrN23fBmDUqlWrXqk2DZoKEySijQAeVXN63C4EBXxeHfGGMMEky2o6vMOnQeRdDmPx8+CK/SBVd2WAApBWo+vNODgRAykeqPk3wTt0AkRWx9PigclN1lTVfnn2fO31dz7+yCj9YibwY7XJ1WDdQ5GCgK0+twzPtfdDvfwGJNa+D2vTYvCRIrARBdsJEAlA00EZWVA79oTW61aoXa+rsQQpNWOVXI5obEnVGZLM3C0WM0a+t2CJd8pLb6Ko+EDDmGCdO28CpA2RkQN9wGjo/QsgjxRDHikCx6pAQgVlZEHkdAL5M2pZlZzi4gmWlVB+de9QvueWGwddmd/14KpVq17v27cvUnGGUrEBwMxptm0XJ+9n/7nyC/a27cXUohv3vXkEJwOjagdInXzpK5lt6wwXw2d+xo2x4vmfrGAlK59Fy278szsfrf/YqNpRbWdjByS3OEMIkSeltBKWhU4d8xDw++p6g6kcCaGc+mqMa12NCSXlDSIiSMtCpw65zo2aaSaF0iWphgCdSQvUxgTnCiFUTVUxb1EDmGCqp9W1C6prEvJuNIdKaF4vVq75EmWHjrDu8egAKgDMTVULpIoJJgDcVXroyB13PTQuMeG5V6Drp8IE/7NFSobP58W2XXvtn905mqa+MmchgN5EtDy5gU2JCVrtclrNf+e9hQYAqKpSjQk6gQ34D1YnSMJ5t41AwCe37tqLXxeM30BEewoKpulNHiFCRDxr3uK253fuqEZj8ZprPyEghHCOr0DTBEqcsRJ0j1ZtDsdjBgX9Plw/5IZWjcUEG6MJRlq2vf2bolJ5x0Pj3CCpXpzX7Ub+autOToatNkmQVAPVtp1qmCY/OXEai1ZXsL/dVdzqouvk/E9WcFU4coKZ5zcmSCrlOEEp5QEhhAqAv961l3oNGA5ihmlZyG7VHHltc1xDpCnC5BruI4hgmCZ27C6CEITK45V4YPhQvP78eEhACoermz5OUAhRPTVFEIgZtpTw6ToOHa5ASem3TRwoCbeNIC0JEo53KKUEkUAg4IMQApASotpgktw4mLMRmKBt2w8DKCwq/ZbHTprB4aowNFVFZTgMj6ZxKJTGaUE/E4CMUBp8ug5VVREKBTkUCrJX90BTFGRkpENVBHxendNDQQ6lBVkIQiDgRzAYgCAFobQAh0JBDgb8TCBkZITg8XigaSrSQ2mcluZnwzQRicbgT0/De/M+5Q8+WsLxmGkCWNAYTLARyht4d/6n7bM6X2f07HcPr/1yi9yxex8/Nm4Ka617cvC8q1nL6cFjJ03nXYXF/O91m7jXDcNZzenOvtyrOPfyAfzBR0t5b1Epvzb7Q27e6Vr2t+vNIqs7D7p7NG/YvIO37ijk4Y+MZzWnBwfPu5p9uVfxn6b9nQv3lfDS1eu46zW3sNa6J+ttruTOVw3hT5au4W+KS3ny9L+bwQ7XcO+B9884Vyq3Trj8mnWbqsPlE4mE7NV/uBsu/6s64fIr1qyXwfZ9Imh2WeVzL75ZN1z+f56rFS5fVB0uf+x4Za1w+d/UC5dfZqs5ParQ4oqq2e/VD5d/ykSo8eHyjSo1CRNGigkT1qqahIl4/YSJhxufMGHOvaFewkSyraj03CZMVENk7udNzHzETXGZ7NxkVafMvOqmsZQz8w01njG3Y+YtbtsqZs6olTJzPzOfcMcbmxyPmTXb5oVun6J6KTNdmHmP27bwP5Iy4746aRc0/x5JUzoz5ya9tHpJU1mnSZrKbSBpKsDMufXndu5vrFJPh/t/nzb3Q+LkD+WH8t9d/g/ovOtFWVwXTQAAAABJRU5ErkJggg=="

CSS_TEMA = r"""
<style>
@import url('https://fonts.googleapis.com/css2?family=Bricolage+Grotesque:opsz,wght@12..96,600;12..96,700;12..96,800&family=Inter:wght@400;500;600&display=swap');
:root{--navy-900:#0B1F3A;--navy-700:#143A6B;--navy-600:#1E4D8C;--navy-100:#DCE6F5;--navy-50:#F3F7FC;--or-700:#C2410C;--or-600:#D9480F;--or-500:#F26B1D;--or-300:#FFB37A;--or-100:#FFE8D6;--or-50:#FFF6EE;--line:#DCE4EF;--display:'Bricolage Grotesque','Segoe UI',system-ui,sans-serif;--texto:'Inter','Segoe UI',system-ui,sans-serif;}
html,body,.stApp,[data-testid="stAppViewContainer"],[data-testid="stMain"]{background:#FFFFFF !important;color:var(--navy-900);}
.stApp,.stApp p,.stApp label,.stApp li,.stApp input,.stApp textarea,.stApp button,.stApp a,.stApp small,.stApp div[data-baseweb="select"]{font-family:var(--texto);}
.stApp p,.stApp label,.stApp li,.stApp .stMarkdown{color:var(--navy-900);}
header[data-testid="stHeader"]{background:transparent;}
[data-testid="stToolbar"],[data-testid="stDecoration"],#MainMenu,footer{visibility:hidden;}
.block-container{max-width:1100px;padding-top:1.2rem;padding-bottom:4rem;}
.ef-top{display:flex;align-items:center;justify-content:space-between;gap:1rem;flex-wrap:wrap;padding-bottom:1rem;border-bottom:3px dotted var(--or-300);}
.ef-logo svg{height:58px;width:auto;display:block;}
.ef-logo .wm{font-family:var(--display);font-weight:800;font-size:35px;letter-spacing:-.6px;}
.ef-logo .tg{font-family:var(--texto);font-weight:500;font-size:10.5px;letter-spacing:.2px;}
.ef-pill{display:inline-flex;align-items:center;gap:.5rem;background:var(--or-100);color:var(--navy-900);font-weight:600;font-size:.86rem;padding:.4rem .85rem;border-radius:999px;}
.ef-pill i{width:.55rem;height:.55rem;border-radius:50%;background:var(--or-500);display:inline-block;}
.ef-hero{padding:2.2rem 0 1.4rem;}
.stApp .ef-hero .ef-h1{font-family:var(--display);font-weight:800;font-size:clamp(2rem,5vw,3.15rem);line-height:1.04;letter-spacing:-.025em;color:var(--navy-900);margin:0 0 .9rem;}
.stApp .ef-hero p{font-size:1.08rem;line-height:1.55;color:#34507A;max-width:58ch;margin:0;}
.st-key-busca{border:4px dotted var(--or-300);border-radius:16px;background:var(--or-50);padding:1rem 1.1rem 1.1rem;}
.st-key-busca div[data-baseweb="input"]{border:2px solid var(--navy-900);border-radius:10px;background:#fff;min-height:3.1rem;}
.st-key-busca div[data-baseweb="base-input"]{background:#fff;}
.st-key-busca div[data-baseweb="input"] input{font-size:1.08rem;color:var(--navy-900);}
.st-key-busca div[data-baseweb="input"]:focus-within{border-color:var(--or-600);box-shadow:0 0 0 3px var(--or-100);}
.st-key-busca .stButton>button{width:100%;min-height:3.1rem;}
.stApp .ef-hint{margin:.7rem 0 0;font-size:.9rem;color:#5A6F92;}
.stApp button[kind="primary"],.stApp button[data-testid="stBaseButton-primary"],.stApp [data-testid="stFormSubmitButton"]>button{background:var(--or-600);border:0;border-radius:10px;font-weight:700;}
.stApp button[kind="primary"] p,.stApp button[data-testid="stBaseButton-primary"] p,.stApp [data-testid="stFormSubmitButton"]>button p{color:#fff;font-weight:700;}
.stApp button[kind="primary"]:hover,.stApp button[data-testid="stBaseButton-primary"]:hover,.stApp [data-testid="stFormSubmitButton"]>button:hover{background:var(--or-700);}
.stApp button[kind="primary"]:disabled,.stApp button[data-testid="stBaseButton-primary"]:disabled{background:var(--or-600);opacity:.45;}
.stApp button[data-testid="stBaseButton-secondary"],.stApp .stDownloadButton>button{background:#fff;border:2px solid var(--navy-900);border-radius:10px;}
.stApp button[data-testid="stBaseButton-secondary"] p,.stApp .stDownloadButton>button p{color:var(--navy-900);font-weight:600;}
.stApp button:focus-visible{outline:3px solid var(--or-500);outline-offset:2px;}
section[data-testid="stSidebar"]{background:var(--navy-50);border-right:1px solid var(--line);}
section[data-testid="stSidebar"] [data-testid="stExpander"]{background:#fff;border:1px solid var(--line);border-radius:12px;}
.ef-side-title{font-family:var(--display);font-weight:800;font-size:1.3rem;color:var(--navy-900);margin:.1rem 0 .8rem;}
span[data-baseweb="tag"]{background:var(--navy-900) !important;color:#fff !important;}
div[data-testid="stAlert"]{background:var(--or-50);border:1px solid var(--or-300);border-radius:12px;color:var(--navy-900);}
div[data-testid="stProgress"] div[role="progressbar"]>div,.stProgress>div>div>div>div{background-color:var(--or-500) !important;}
.ef-resumo{display:flex;align-items:baseline;gap:.8rem;flex-wrap:wrap;margin:1.6rem 0 .3rem;}
.ef-resumo b{font-family:var(--display);font-weight:800;font-size:2.5rem;line-height:1;color:var(--or-600);}
.ef-resumo span{font-size:1.05rem;color:var(--navy-900);}
.ef-chips{display:flex;flex-wrap:wrap;gap:.45rem;margin:1.2rem 0 0;}
.ef-chip{font-size:.82rem;font-weight:600;padding:.25rem .7rem;border-radius:999px;border:1px solid var(--navy-100);background:var(--navy-50);color:var(--navy-700);}
.ef-chip b{margin-left:.25rem;}
.ef-chip.ok{background:var(--navy-900);border-color:var(--navy-900);color:#fff;}
.ef-chip.err{background:var(--or-100);border-color:var(--or-300);color:var(--or-700);}
.ef-grupo{margin:1.6rem 0 0;}
.ef-grupo-h{display:flex;align-items:center;justify-content:space-between;gap:1rem;padding-bottom:.55rem;border-bottom:2px solid var(--navy-900);}
.ef-grupo-h .ef-h3{font-family:var(--display);font-weight:700;font-size:1.2rem;color:var(--navy-900);}
.ef-badge{background:var(--navy-900);color:#fff;font-weight:600;font-size:.8rem;padding:.2rem .65rem;border-radius:999px;white-space:nowrap;}
.stApp ul.ef-lotes{list-style:none !important;margin:0 !important;padding:0 !important;}
.stApp ul.ef-lotes li{display:flex;align-items:flex-start;justify-content:space-between;gap:1rem;margin:0;padding:.8rem 0;border-bottom:2px dotted var(--or-300);}
.stApp ul.ef-lotes a.t{color:var(--navy-700);font-weight:600;line-height:1.4;text-decoration:none;display:-webkit-box;-webkit-line-clamp:2;-webkit-box-orient:vertical;overflow:hidden;}
.stApp ul.ef-lotes a.t:hover{color:var(--or-600);text-decoration:underline;}
.stApp ul.ef-lotes a.a{flex:none;font-size:.8rem;color:var(--or-700);border:1px solid var(--or-300);border-radius:999px;padding:.15rem .65rem;text-decoration:none;white-space:nowrap;}
.stApp ul.ef-lotes a.a:hover{background:var(--or-100);}
.stApp ul.ef-lotes a:focus-visible{outline:3px solid var(--or-500);outline-offset:2px;border-radius:4px;}
.stApp ul.ef-links{margin:.2rem 0 .8rem;padding-left:1.2rem;}
.stApp ul.ef-links a{color:var(--navy-700);font-weight:500;}
.stApp .ef-sub{font-family:var(--display);font-weight:700;font-size:1rem;margin:1.2rem 0 .2rem;color:var(--navy-900);}
.ef-foot{margin-top:3rem;padding-top:1rem;border-top:3px dotted var(--or-300);color:#5A6F92;font-size:.85rem;}
@media (max-width:640px){.stApp ul.ef-lotes li{flex-direction:column;gap:.4rem;}.ef-logo svg{height:48px;}}
@media (prefers-reduced-motion:reduce){*{transition:none !important;animation:none !important;}}
</style>
"""


def _icone_pagina():
    try:
        from PIL import Image
        return Image.open(io.BytesIO(base64.b64decode(ICONE_PNG_B64)))
    except Exception:
        return "📮"


st.set_page_config(
    page_title="eFilatelia · busca de selos em leilões",
    page_icon=_icone_pagina(),
    layout="wide",
)
st.markdown(CSS_TEMA, unsafe_allow_html=True)
st.markdown(
    '<div class="ef-top"><div class="ef-logo" role="img" aria-label="eFilatelia">'
    + LOGO_SVG
    + '</div><span class="ef-pill"><i></i>Leilões em andamento</span></div>'
    '<div class="ef-hero"><div class="ef-h1" role="heading" aria-level="1">'
    "Todos os leilões de selos<br>em uma só busca</div>"
    "<p>Digite o selo que você procura. O eFilatelia consulta os leilões abertos agora "
    "e mostra cada lote com link direto para o leiloeiro.</p></div>",
    unsafe_allow_html=True,
)

if "sites_v20" not in st.session_state:
    st.session_state.sites_v20 = list(SITES_PADRAO) + carregar_extras()

with st.sidebar:
    st.markdown('<div class="ef-side-title">Sites e filtros</div>', unsafe_allow_html=True)
    nomes = [s.nome for s in st.session_state.sites_v20]
    padrao = [s.nome for s in st.session_state.sites_v20 if s.ativo]

    with st.expander("Sites e limites", expanded=True):
        escolhidos = st.multiselect("Sites a consultar", nomes, default=padrao, key="sel_sites")
        if st.button("Restaurar sites padrão", key="restaurar_sites"):
            st.session_state.pop("sel_sites", None)
            st.session_state.sites_v20 = list(SITES_PADRAO) + carregar_extras()
            st.rerun()
        max_paginas = st.slider("Páginas de resultados por site", 1, 30, 5)
        limite_s = st.slider("Tempo máximo por site (segundos)", 15, 180, 45)

    with st.expander("Filtros", expanded=True):
        exigir_titulo = st.checkbox(
            "Exigir todos os termos no título do lote", value=True,
            help="Ligado (recomendado): descarta lotes cujo título/descrição não contenha todas as "
                 "palavras buscadas. Alguns sites preenchem a lista com itens sem relação.",
        )
        rigor = st.radio(
            "Filtro de selos",
            ["Desligado", "Moderado (recomendado)", "Estrito"],
            index=1,
            help="Moderado: descarta títulos que parecem moeda, cédula, relógio, joia, boneco, "
                 "álbum, livro etc. (a menos que citem selo/filatelia). "
                 "Estrito: só aceita títulos com palavras de selo (selo, RHM, Yvert, Scott, "
                 "bloco, série, carimbo...). Desligado: não filtra.",
        )
        verificar = st.checkbox(
            "Conferir na página do lote se não foi vendido", value=False,
            help="Mais lento: abre cada lote e descarta os marcados como vendido/encerrado.",
        )

    with st.expander("Avançado", expanded=False):
        usar_navegador = st.checkbox(
            "Usar navegador automatizado (Playwright) nos sites com JavaScript", value=False,
            help="Necessário para sites que montam os resultados com JavaScript (ex.: Philasearch, "
                 "Catawiki). Requer: pip install playwright && playwright install chromium. Mais lento.",
        )
        respeitar = st.checkbox("Respeitar robots.txt", value=True)

    with st.expander("Adicionar site", expanded=False):
        st.caption("Abra qualquer site de leilão, faça uma busca por um termo e cole aqui a URL "
                   "da página de resultados. O eFilatelia troca o termo por {q} e a página por {pag}.")
        if st.session_state.get("msg_novo"):
            st.success(st.session_state.pop("msg_novo"))
        with st.form("novo_url", clear_on_submit=True):
            u_nome = st.text_input("Nome do site")
            u_url = st.text_input("URL da página de resultados", placeholder="https://site.com/lots?q=guiana")
            u_termo = st.text_input("Termo que você buscou nessa URL", placeholder="guiana")
            u_js = st.checkbox("Site precisa de JavaScript (Playwright)", value=False)
            if st.form_submit_button("Adicionar site") and u_nome and u_url and u_termo:
                novo, explicacao = site_a_partir_de_url(u_nome, u_url, u_termo, u_js)
                st.session_state.sites_v20 = [
                    x for x in st.session_state.sites_v20 if x.nome != novo.nome
                ] + [novo]
                salvar_extras(st.session_state.sites_v20)
                st.session_state.msg_novo = explicacao
                st.rerun()

        extras_atuais = [x.nome for x in st.session_state.sites_v20 if x.extra]
        if extras_atuais:
            rem = st.selectbox("Remover site adicionado", ["—"] + extras_atuais)
            if rem != "—" and st.button("Remover site"):
                st.session_state.sites_v20 = [x for x in st.session_state.sites_v20 if x.nome != rem]
                salvar_extras(st.session_state.sites_v20)
                st.rerun()

        if st.checkbox("Informar seletor CSS manualmente"):
            with st.form("novo", clear_on_submit=True):
                n_nome = st.text_input("Nome")
                n_url = st.text_input("URL de busca (use {q})", placeholder="https://site.com/busca?q={q}")
                n_sel = st.text_input("Seletor CSS dos links dos lotes", value='a[href*="lote"]')
                n_enc = st.selectbox("Codificação da busca", ["utf-8", "latin-1"])
                n_js = st.checkbox("Site precisa de JavaScript (Playwright)", value=False)
                if st.form_submit_button("Adicionar site") and n_nome and "{q}" in n_url:
                    st.session_state.sites_v20.append(
                        Site(n_nome, n_url, n_sel, n_enc, usa_paginas="{pag}" in n_url,
                             js=n_js, extra=True)
                    )
                    salvar_extras(st.session_state.sites_v20)
                    st.rerun()

# ---- Caixa de busca (moldura pontilhada, como a serrilha de um selo) ----
try:
    caixa = st.container(key="busca")
except TypeError:  # Streamlit antigo, sem 'key' em container
    caixa = st.container()
with caixa:
    col_txt, col_btn = st.columns([5, 1.4])
    with col_txt:
        consulta = st.text_input(
            "Selo buscado",
            placeholder="Ex.: Olho de Boi, Penny Black, RHM B-45",
            label_visibility="collapsed",
        )
    with col_btn:
        buscar = st.button("Buscar selos", type="primary", disabled=not consulta.strip())
st.markdown(
    f'<p class="ef-hint">{len(escolhidos)} sites selecionados. '
    "Ajuste sites e filtros no menu lateral.</p>",
    unsafe_allow_html=True,
)


def _link_direto_ok(url: str):
    """True se a página do lote abre sem erro; False se der erro HTTP/erro do IIS; None se não deu para testar."""
    try:
        r = requests.get(url, headers=HEADERS, timeout=(5, 10))
    except requests.RequestException:
        return None
    if r.status_code >= 400 or "custom error module" in r.text.lower():
        return False
    return True


def marcar_links_diretos(lotes: list[dict], limite_s_: int = 25) -> None:
    """
    Nos lotes do portal LeilõesBR, o link direto (página do lote na casa de leilões) é o
    principal. Antes de mostrar, testa cada um; se der erro, o link do portal passa a ser
    o principal. Modifica a lista no lugar (campo 'ok_direto').
    """
    alvo = [l for l in lotes if l["Link"] != l["Link do portal"]][:60]
    if not alvo:
        return
    ex = ThreadPoolExecutor(max_workers=8)
    futs = {ex.submit(_link_direto_ok, l["Link"]): l for l in alvo}
    try:
        for f in as_completed(futs, timeout=limite_s_):
            futs[f]["ok_direto"] = f.result()
    except FutTimeout:
        pass
    finally:
        ex.shutdown(wait=False, cancel_futures=True)


def executar_busca(sites_sel, consulta_):
    """Consulta os sites em paralelo e devolve (lotes, estatísticas por site)."""
    todos_, stats_ = [], []
    nomes_sel = [s_.nome for s_ in sites_sel]
    barra = st.progress(0.0, text="Consultando: " + ", ".join(nomes_sel)[:200])
    ex = ThreadPoolExecutor(max_workers=max(1, min(8, len(sites_sel))))
    futs = {
        ex.submit(processar_site, s_, consulta_, max_paginas, respeitar,
                  exigir_titulo, verificar, rigor, usar_navegador, limite_s): s_
        for s_ in sites_sel
    }
    pendentes = set(nomes_sel)
    concluidos = 0
    try:
        for f in as_completed(futs, timeout=limite_s + 45):
            s_ = futs[f]
            try:
                res, stt = f.result()
            except Exception as e:  # nenhum site deve derrubar a busca inteira
                res, stt = [], novo_stats(s_.nome)
                stt["Erro"] = f"Erro inesperado: {type(e).__name__}: {e}"
            todos_.extend(res)
            stats_.append(stt)
            pendentes.discard(s_.nome)
            concluidos += 1
            aguardando = ", ".join(sorted(pendentes))[:200] or "nenhum"
            barra.progress(concluidos / len(futs),
                           text=f"{concluidos} de {len(futs)} sites concluídos. Aguardando: {aguardando}")
    except FutTimeout:
        for nome in sorted(pendentes):
            stt = novo_stats(nome)
            stt["Erro"] = "Tempo esgotado: o site não respondeu a tempo (ignorado nesta busca)"
            stats_.append(stt)
    finally:
        ex.shutdown(wait=False, cancel_futures=True)  # não espera threads presas
    barra.text = "Conferindo os links dos lotes..."
    marcar_links_diretos(todos_)
    barra.empty()
    return todos_, stats_


if buscar:
    sites_escolhidos = [s for s in st.session_state.sites_v20 if s.nome in escolhidos]
    if not sites_escolhidos:
        sites_escolhidos = [s for s in st.session_state.sites_v20 if s.ativo]
        st.info("Nenhum site estava marcado no menu lateral; usei todos os sites padrão.")
    if True:
        todos_b, stats_b = executar_busca(sites_escolhidos, consulta.strip())
        # guarda o resultado: clicar em download ou trocar um seletor não apaga a lista
        st.session_state.ultimo = {
            "consulta": consulta.strip(), "todos": todos_b,
            "stats": stats_b, "sites": sites_escolhidos,
        }


def _lista_links(itens, consulta_):
    linhas = []
    for sx, stt in itens:
        url_ = html.escape(url_busca_manual(sx, consulta_), quote=True)
        motivo = f" — <small>{html.escape(stt['Erro'][:110])}</small>" if stt["Erro"] else ""
        linhas.append(f'<li><a href="{url_}" target="_blank" rel="noopener">'
                      f"{html.escape(sx.nome)}</a>{motivo}</li>")
    return '<ul class="ef-links">' + "".join(linhas) + "</ul>"


ultimo = st.session_state.get("ultimo")
if ultimo:
    q = ultimo["consulta"]
    todos, stats_list, sites_sel = ultimo["todos"], ultimo["stats"], ultimo["sites"]

    _ordem = {sx.nome: k for k, sx in enumerate(sites_sel)}
    _chips = []
    for stt in sorted(stats_list, key=lambda x: _ordem.get(x["Site"], 99)):
        n_ = stt["Mantidos"]
        classe = "ok" if n_ else ("err" if stt["Erro"] else "zero")
        dica = stt["Erro"] or stt.get("Nota") or ("nenhum lote para este termo" if not n_ else "")
        _chips.append(f'<span class="ef-chip {classe}" title="{html.escape(dica, quote=True)}">'
                      f'{html.escape(stt["Site"].split(" (")[0])} <b>{n_}</b></span>')
    st.markdown('<div class="ef-chips">' + "".join(_chips) + "</div>", unsafe_allow_html=True)

    if todos:
        df = pd.DataFrame(todos).drop_duplicates(subset="Link")
        n = len(df)
        st.markdown(
            f'<div class="ef-resumo"><b>{n}</b><span>{"lote" if n == 1 else "lotes"} em leilões '
            f"em andamento para “{html.escape(q)}”</span></div>",
            unsafe_allow_html=True,
        )
        for nome_site, g in df.groupby("Site", sort=False):
            itens_html = []
            for _, row in g.iterrows():
                t = html.escape(row["Título"])
                principal, secundario, rotulo = links_do_lote(row)
                a_princ = html.escape(principal, quote=True)
                alt_html = ""
                if secundario:
                    alt_html = (f'<a class="a" href="{html.escape(secundario, quote=True)}" '
                                f'target="_blank" rel="noopener">{rotulo}</a>')
                itens_html.append(
                    f'<li><a class="t" href="{a_princ}" target="_blank" rel="noopener">{t}</a>'
                    f"{alt_html}</li>"
                )
            st.markdown(
                '<div class="ef-grupo"><div class="ef-grupo-h">'
                f'<div class="ef-h3" role="heading" aria-level="2">{html.escape(nome_site)}</div>'
                f'<span class="ef-badge">{len(g)} {"lote" if len(g) == 1 else "lotes"}</span></div>'
                '<ul class="ef-lotes">' + "".join(itens_html) + "</ul></div>",
                unsafe_allow_html=True,
            )
        st.markdown('<div class="ef-sub">Baixar resultados</div>', unsafe_allow_html=True)
        base_nome = f"efilatelia_{_slug(q)}_{datetime.now():%Y%m%d}"
        c_xlsx, c_docx, c_csv = st.columns(3)
        try:
            c_xlsx.download_button(
                "Excel (.xlsx)", gerar_excel(df, q, stats_list), key="dl_xlsx",
                file_name=f"{base_nome}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        except ImportError:
            c_xlsx.info("Para baixar em Excel, rode: pip install openpyxl")
        try:
            c_docx.download_button(
                "Word (.docx)", gerar_word(df, q, stats_list), key="dl_docx",
                file_name=f"{base_nome}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        except ImportError:
            c_docx.info("Para baixar em Word, rode: pip install python-docx")
        c_csv.download_button(
            "CSV", df.to_csv(index=False).encode("utf-8-sig"), key="dl_csv",
            file_name=f"{base_nome}.csv", mime="text/csv",
        )
    else:
        st.warning(f"Nenhum lote encontrado para “{q}”. O que aconteceu em cada site:")
        for stt in stats_list:
            if stt["Erro"]:
                motivo = stt["Erro"]
            elif stt["Links achados"] == 0:
                motivo = (f"HTTP {stt['HTTP']}, mas nenhum link de lote reconhecido "
                          "(seletor CSS não bate ou a página é montada por JavaScript — "
                          "tente ativar o Playwright em Avançado).")
            else:
                motivo = (f"{stt['Links achados']} link(s) achado(s), todos descartados pelos filtros "
                          f"(título: {stt['Descartados (título)']}, "
                          f"não é selo: {stt['Descartados (não é selo)']}, "
                          f"encerrado: {stt['Descartados (encerrado)']}).")
            st.markdown(f"- **{stt['Site']}**: {motivo}")

    por_nome = {stt["Site"]: stt for stt in stats_list}
    bloqueados = [(sx, por_nome[sx.nome]) for sx in sites_sel
                  if sx.nome in por_nome and por_nome[sx.nome]["Erro"]]
    vazios = [(sx, por_nome[sx.nome]) for sx in sites_sel
              if sx.nome in por_nome and not por_nome[sx.nome]["Erro"]
              and por_nome[sx.nome]["Links achados"] == 0]
    if bloqueados:
        st.markdown('<div class="ef-sub">Sites que o eFilatelia não conseguiu ler. '
                    "Abra a busca direto no navegador:</div>", unsafe_allow_html=True)
        st.markdown(_lista_links(bloqueados, q), unsafe_allow_html=True)
    if vazios:
        st.markdown('<div class="ef-sub">Sites que responderam sem nenhum lote reconhecido '
                    "(pode não haver lotes para o termo, ou a página depende de JavaScript):</div>",
                    unsafe_allow_html=True)
        st.markdown(_lista_links(vazios, q), unsafe_allow_html=True)

    if todos:
        with st.expander("Testar os links (mostra o que o servidor responde)"):
            testar = st.button("Testar os 8 primeiros links", key="testar_links")
            if testar:
                amostra = pd.DataFrame(todos).drop_duplicates(subset="Link").head(8)
                linhas_teste = []
                for _, row in amostra.iterrows():
                    try:
                        r = requests.get(row["Link"], headers=HEADERS, timeout=TIMEOUT)
                        linhas_teste.append({
                            "Título": row["Título"][:60],
                            "HTTP": str(r.status_code),
                            "URL final": r.url,
                            "Erro do IIS?": "sim" if "custom error module" in r.text.lower() else "não",
                        })
                    except requests.RequestException as e:
                        linhas_teste.append({"Título": row["Título"][:60], "HTTP": "falha",
                                             "URL final": str(e)[:120], "Erro do IIS?": "?"})
                st.dataframe(pd.DataFrame(linhas_teste), hide_index=True)

    with st.expander("Diagnóstico por site", expanded=not todos):
        diag_df = pd.DataFrame(stats_list).drop(columns=["_html", "_amostra"])
        st.dataframe(diag_df, hide_index=True)
        st.caption(
            "HTTP diferente de 200 ou Erro preenchido: o site bloqueou ou recusou. "
            "Links achados igual a 0 com HTTP 200: o seletor CSS não bate com o HTML "
            "ou o conteúdo é carregado por JavaScript."
        )
        alvo = st.selectbox("Ver detalhes de um site", [x["Site"] for x in stats_list])
        det = next(x for x in stats_list if x["Site"] == alvo)
        st.code(det["URL"])
        if det["_amostra"]:
            st.markdown("**Exemplos de títulos descartados pelo filtro de título:**")
            st.markdown("\n".join(f"- {t}" for t in det["_amostra"]))
        st.code(det["_html"] or "(vazio)", language="html")

st.markdown(
    '<div class="ef-foot">O eFilatelia reúne resultados de sites de leilão e leva você ao lote '
    "original. Lances e compras acontecem no site do leiloeiro.</div>",
    unsafe_allow_html=True,
)
